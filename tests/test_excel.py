"""Excel: Kompatibilität mit bestehenden Dateien und verlustfreier Round-Trip.

Aufruf: python -m unittest discover -s tests
"""
import io
import json
import os
import unittest

import openpyxl

from app.excel import schema as S
from app.excel.reader import load_gradebook
from app.excel.writer import build_gradebook
from app.grades import berechnung as B

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEMO = os.path.join(ROOT, "Noten_demo_klasse.xlsx")
REFERENZ = os.path.join(ROOT, "tests", "fixtures", "referenz_demo_klasse.json")


def _load_demo():
    with open(DEMO, "rb") as f:
        raw = f.read()
    try:
        return load_gradebook(raw, None)
    except Exception:
        return load_gradebook(raw, "test")


def _proposals(d):
    out = {}
    for s in d["stammdaten"]:
        n = f"{s['nachname']}, {s['vorname']}"
        r = {sl: B.sl_note_for(d, n, sl) for sl in B.SL_KEYS}
        r["HJ1"] = B.hj_vorschlag_for(d, n, "HJ1")
        r["HJ2"] = B.hj_vorschlag_for(d, n, "HJ2")
        r["SJ"] = B.compute_schuljahr_note_klasse(
            n, d.get("hj_noten") or {}, s.get("aufnahme_ab_hj"),
            s.get("aufnahme_vorherige_noten"), B.get_gewichtung(d))
        out[n] = r
    return out


@unittest.skipUnless(os.path.exists(DEMO), "Demo-Datei fehlt")
class AltdateiTest(unittest.TestCase):
    """Bestehende Dateien müssen einlesbar sein und dieselben Vorschläge liefern."""

    def _assert_referenz(self, d):
        with open(REFERENZ, encoding="utf-8") as f:
            ref = json.load(f)["students"]
        got = _proposals(d)
        self.assertEqual(set(ref), set(got))
        for name, werte in ref.items():
            for key, exp in werte.items():
                if exp is None:
                    self.assertIsNone(got[name][key], (name, key))
                else:
                    self.assertAlmostEqual(got[name][key], exp, places=5, msg=(name, key))

    def test_demo_wie_vorher(self):
        self._assert_referenz(_load_demo())

    def test_demo_nach_roundtrip_wie_vorher(self):
        d = load_gradebook(build_gradebook(_load_demo(), "geheim"), "geheim")
        self._assert_referenz(d)


class RoundTripTest(unittest.TestCase):
    def _data(self):
        name = "Muster, Max"
        return {
            "klasse": "7p", "fach": "Mathe", "schuljahr": "2526", "modus": "klasse",
            "stammdaten": [{"nachname": "Muster", "vorname": "Max", "status": "Aktiv"}],
            "leistungsnachweise": [
                {"sheet_name": "LN_Test1", "name": "Test1", "ln_typ": "KLN", "sl_zuordnung": "SL1",
                 "hj": None, "nachtermin_von": None, "noten_runden": True,
                 "aufgaben": [{"label": "1", "afb": "I", "max_punkte": 10}],
                 "schueler": [{"name": name, "punkte": [7], "note_15": None}]},
                {"sheet_name": "LN_Referat", "name": "Referat", "ln_typ": "MDL", "sl_zuordnung": "SL1",
                 "hj": None, "nachtermin_von": None, "noten_runden": False, "aufgaben": [],
                 "thema": "Vulkane", "datum": "2025-05-25",
                 "schueler": [{"name": name, "punkte": [], "note_15": 13, "ignoriert": False}]},
                {"sheet_name": "LN_KA1", "name": "KA1", "ln_typ": "GLN", "sl_zuordnung": None,
                 "hj": "HJ1", "nachtermin_von": None, "noten_runden": True,
                 "aufgaben": [{"label": "1", "afb": "I", "max_punkte": 20}],
                 "schueler": [{"name": name, "punkte": [15], "note_15": None}]},
            ],
            "mdl_noten": {name: {"SL1": 12}},
            "sl_noten_actual": {}, "hj_noten": {}, "schuljahr_noten_actual": {},
            "sl_gewichtung": {"sl_mdl_pct": 60.0, "sl_kln_pct": 40.0, "sl_mittel_w": 2.0,
                              "sj_hj1_w": 1.0, "sj_hj2_w": 1.0},
            "kln_weights": {"SL1": {"LN_Test1": 2.0}},
            "mdl_weights": {"SL1": {"LN_Referat": 3.0}},
            "gln_weights": {"HJ1": {"LN_KA1": 1.5}},
        }

    def test_alles_bleibt_erhalten(self):
        src = self._data()
        d = load_gradebook(build_gradebook(src, "pw"), "pw")
        for key in ("sl_gewichtung",):
            for k, v in src[key].items():
                self.assertEqual(d[key][k], v, k)
        for key in ("kln_weights", "mdl_weights", "gln_weights"):
            self.assertEqual(d[key], src[key], key)
        self.assertEqual(d["mdl_noten"]["Muster, Max"]["SL1"], 12)
        ref = next(ln for ln in d["leistungsnachweise"] if ln["sheet_name"] == "LN_Referat")
        self.assertEqual(ref["ln_typ"], "MDL")
        self.assertEqual(ref["sl_zuordnung"], "SL1")
        self.assertEqual(ref["schueler"][0]["note_15"], 13)
        self.assertEqual((ref["thema"], ref["datum"]), ("Vulkane", "2025-05-25"))
        test1 = next(ln for ln in d["leistungsnachweise"] if ln["sheet_name"] == "LN_Test1")
        self.assertEqual((test1["thema"], test1["datum"]), ("", ""))
        self.assertEqual(B.mdl_for(d, "Muster, Max", "SL1"), (12.0, "fest"))

    def test_einstellungen_sind_lesbar_beschriftet(self):
        raw = build_gradebook(self._data(), None)
        ws = openpyxl.load_workbook(io.BytesIO(raw))[S.SHEET_EINSTELLUNGEN]
        texte = [ws.cell(r, S.ES_COL_INFO).value for r in range(2, ws.max_row + 1)]
        self.assertIn('Gewicht MDL "Referat" in SL1', texte)

    def test_datum_in_excel_von_hand_geaendert(self):
        raw = build_gradebook(self._data(), None)
        wb = openpyxl.load_workbook(io.BytesIO(raw))
        wb["LN_Referat"].cell(S.LN_ROW_META, S.LN_META_DATUM_VAL).value = "3.6.2025"
        buf = io.BytesIO(); wb.save(buf)
        d = load_gradebook(buf.getvalue(), None)
        ref = next(ln for ln in d["leistungsnachweise"] if ln["sheet_name"] == "LN_Referat")
        self.assertEqual(ref["datum"], "2025-06-03")

    def test_gewicht_geloeschter_ln_wird_nicht_geschrieben(self):
        src = self._data()
        src["kln_weights"]["SL1"]["LN_Weg"] = 5.0
        d = load_gradebook(build_gradebook(src, None), None)
        self.assertNotIn("LN_Weg", d["kln_weights"]["SL1"])


if __name__ == "__main__":
    unittest.main()
